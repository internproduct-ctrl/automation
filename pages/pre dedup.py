import streamlit as st
import pandas as pd
import io
import sys, os
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from theme import inject, page_header, metric_cards
from openpyxl import load_workbook
from openpyxl.styles import PatternFill
from datetime import datetime
import time

st.set_page_config(page_title="Pre Dedup", page_icon="📋", layout="wide")
inject()

def get_manager_column(df, base):
    base_lower = base.strip().lower()

    for col in df.columns:
        if str(col).strip().lower() == base_lower:
            return col

    return None
# ── HELPERS ───────────────────────────────────────────────────────────────────
def normalize(val):
    if pd.isna(val):
        return None
    return str(val).strip().replace(" ", "").upper()

def check_match(values):
    vals = [normalize(v) for v in values if normalize(v) not in [None, ""]]
    if len(vals) == 0: return False
    if len(vals) == 1: return True
    return len(set(vals)) == 1

def to_bool(val):
    if isinstance(val, bool):
        return val

    if pd.isna(val):
        return False

    val = str(val).strip().upper()
    val = val.replace("\n", "").replace("\r", "")

    if val in ["TRUE", "T", "YES", "Y", "1"]:
        return True

    if val in ["FALSE", "F", "NO", "N", "0", ""]:
        return False

    return False

def style_table(df):
    styles = pd.DataFrame("", index=df.index, columns=df.columns)

    for idx in df.index:
        ok = df.at[idx, "Overall_Status"]

        if ok == "✔ MATCH":
            bg = "#ecfdf5"      # light green
            text_c = "#065f46"  # dark green text
            accent = "#16a34a"  # highlight
        else:
            bg = "#fef2f2"      # light red
            text_c = "#7f1d1d"  # dark red text
            accent = "#dc2626"

        for col in df.columns:
            if col == "Overall_Status":
                styles.at[idx, col] = (
                    f"background-color:{bg};"
                    f"color:{accent};"
                    f"font-weight:700;"
                    f"border:1px solid #e5e7eb;"
                )
            else:
                styles.at[idx, col] = (
                    f"background-color:{bg};"
                    f"color:{text_c};"
                    f"border:1px solid #e5e7eb;"
                )
    return styles

# ── UI ────────────────────────────────────────────────────────────────────────
page_header("PRE DEDUP", "Excel Column-Group Duplicate Validator")

uploaded_file = st.file_uploader(
    "Upload Excel File (.xlsx)", type=["xlsx"],
    label_visibility="collapsed"
)

if uploaded_file is None:
    st.markdown(
        '<div class="empty-state">'
        '<div class="empty-icon">📋</div>'
        '<div class="empty-txt">Upload an Excel file to begin validation</div>'
        '</div>',
        unsafe_allow_html=True
    )
    st.stop()

# ── READ FILE ─────────────────────────────────────────────────────────────────
file_bytes = uploaded_file.read()
df = pd.read_excel(io.BytesIO(file_bytes))

# ── DETECT COLUMN GROUPS ──────────────────────────────────────────────────────
groups = {}
for col in df.columns:
    parts = col.rsplit(" ", 1)
    if len(parts) == 2 and parts[1].isdigit():
        base = parts[0].strip()
        groups.setdefault(base, []).append(col)

if not groups:
    st.error("No grouped columns found (e.g. 'Name 1', 'Name 2').")
    st.stop()

st.markdown(
    '<div style="font-family:\'IBM Plex Mono\',monospace;font-size:0.68rem;'
    'letter-spacing:2px;text-transform:uppercase;color:#3d5070;margin-bottom:8px;">'
    'Select column groups to validate</div>',
    unsafe_allow_html=True
)
selected_bases = st.multiselect("", list(groups.keys()), label_visibility="collapsed")

if not selected_bases:
    st.markdown(
        '<div class="cpill">Choose one or more column groups above to run validation</div>',
        unsafe_allow_html=True
    )
    st.stop()

# ── COMPUTE ───────────────────────────────────────────────────────────────────
for base in selected_bases:
    df[f"{base}_Calc"] = df[groups[base]].apply(
        lambda row: check_match(row.values), axis=1
    )

def overall_check(row):
    for base in selected_bases:

        # 🔥 FIX: get correct column (case insensitive)
        manager_col = get_manager_column(df, base)

        if manager_col is None:
            return False

        manager_val = to_bool(row[manager_col])

        # 🔥 FIX: force boolean
        calc_val = True if row[f"{base}_Calc"] else False

        # 🔥 FIX: compare same type
        if bool(manager_val) != bool(calc_val):
            return False

    return True

df["Overall_Status"] = df.apply(overall_check, axis=1)

# ── METRICS ───────────────────────────────────────────────────────────────────
total   = len(df)
correct = int(df["Overall_Status"].sum())
wrong   = total - correct

metric_cards([
    ("mc-total", total,   "Total Rows"),
    ("mc-new",   correct, "Correct"),
    ("mc-imp",   wrong,   "Mismatched"),
])

# ── FILTER ────────────────────────────────────────────────────────────────────
f1, f2, _ = st.columns([2, 2, 5])
sel_status = f1.selectbox("STATUS", ["ALL", "✔ MATCH", "✘ MISMATCH"], key="pd_status")
sel_base   = f2.selectbox("GROUP",  ["ALL"] + selected_bases,          key="pd_base")

# Build display df with readable status
display = df.copy()
display["Overall_Status"] = display["Overall_Status"].map(
    {True: "✔ MATCH", False: "✘ MISMATCH"}
)

# Apply filters
filt = display.copy()
if sel_status != "ALL":
    filt = filt[filt["Overall_Status"] == sel_status]
if sel_base != "ALL":
    keep = groups[sel_base] + [f"{sel_base}_Calc", "Overall_Status"]
    keep = [c for c in keep if c in filt.columns]
    filt = filt[keep]

st.markdown(
    f'<div class="cpill">Showing {len(filt)} of {total} rows</div>',
    unsafe_allow_html=True
)

# ── TABLE ─────────────────────────────────────────────────────────────────────
if filt.empty:
    st.info("No records match the selected filters.")
else:
    st.subheader("📊 Results")

    # placeholder for dynamic updates
    table_placeholder = st.empty()

    chunk_size = 200
    total_rows = len(filt)

    for i in range(0, total_rows, chunk_size):
        partial_df = filt.iloc[:i + chunk_size]

        table_placeholder.dataframe(
            partial_df.style.apply(style_table, axis=None),
            use_container_width=True,
            height=600
        )

        time.sleep(0.05)

# ── DOWNLOAD ──────────────────────────────────────────────────────────────────
st.markdown("<br>", unsafe_allow_html=True)

output = io.BytesIO()
wb = load_workbook(io.BytesIO(file_bytes))
ws = wb.active

date_format_map = {}
for col_idx, col_name in enumerate(df.columns, start=1):
    cell = ws.cell(row=2, column=col_idx)
    if isinstance(cell.value, datetime):
        date_format_map[col_name] = cell.number_format

with pd.ExcelWriter(output, engine='openpyxl') as writer:
    df.to_excel(writer, index=False)
    ws_out = writer.sheets["Sheet1"]
    red   = PatternFill(start_color="FF4D4D", fill_type="solid")
    green = PatternFill(start_color="22C55E", fill_type="solid")
    for i, row in df.iterrows():
        excel_row = i + 2
        fill = green if row["Overall_Status"] else red
        for col in range(1, len(df.columns) + 1):
            ws_out.cell(excel_row, col).fill = fill
        for j, col_name in enumerate(df.columns, start=1):
            if col_name in date_format_map:
                ws_out.cell(excel_row, j).number_format = date_format_map[col_name]

output.seek(0)
st.download_button(
    "📥  Download Results (Excel)",
    output,
    "pre_dedup_output.xlsx",
    use_container_width=False,
)